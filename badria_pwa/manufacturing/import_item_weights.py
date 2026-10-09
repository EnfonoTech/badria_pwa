# badria_pwa.manufacturing.import_item_weights - One-off loader for the
# finished goods weight sheet ("product name.xlsx": CODE | FINAL PRODUCTS |
# QTY | UOM | WT (kg per box) | PCS QTY | PCS Wt (kg per piece)).
#
# Production costing splits material cost by kg, from Item.weight_per_unit
# (kg per stock unit = per piece). Run once per site, e.g.:
#   bench --site <site> execute badria_pwa.manufacturing.import_item_weights.run \
#       --kwargs "{'path': '/path/to/product name.xlsx'}"
# Add 'dry_run': 1 to only print what would change.

import frappe
from frappe.utils import flt

CODE_HEADER = "CODE"
PIECE_WEIGHT_HEADER = "PCS WT"
BOX_WEIGHT_HEADER = "WT"


def run(path, dry_run=0):
	from openpyxl import load_workbook

	sheet = load_workbook(path, data_only=True, read_only=True).worksheets[0]
	rows = list(sheet.iter_rows(values_only=True))
	# The header is not necessarily the first row (the sheet has blank rows on top).
	header_idx = next(
		i for i, row in enumerate(rows) if CODE_HEADER in [str(c or "").strip().upper() for c in row]
	)
	header = [str(c or "").strip().upper() for c in rows[header_idx]]
	code_col = header.index(CODE_HEADER)
	piece_col = header.index(PIECE_WEIGHT_HEADER)
	box_col = header.index(BOX_WEIGHT_HEADER)

	updated, not_found, mismatched = [], [], []
	for row in rows[header_idx + 1 :]:
		code = str(row[code_col] or "").strip()
		piece_kg = flt(row[piece_col])
		if not code or not piece_kg:
			continue

		if not frappe.db.exists("Item", code):
			not_found.append(code)
			continue

		# Sanity check: box kg / piece kg should be the Carton conversion.
		carton_factor = flt(
			frappe.db.get_value("UOM Conversion Detail", {"parent": code, "uom": "Carton"}, "conversion_factor")
		)
		box_kg = flt(row[box_col])
		if carton_factor and box_kg and abs(box_kg / piece_kg - carton_factor) > 0.01:
			mismatched.append(f"{code}: box {box_kg} kg / piece {piece_kg} kg != Carton {carton_factor}")

		updated.append(f"{code}: {piece_kg} kg per piece")
		if not int(dry_run):
			frappe.db.set_value("Item", code, {"weight_per_unit": piece_kg, "weight_uom": "Kg"})

	if not int(dry_run):
		frappe.db.commit()

	print(("Would update" if int(dry_run) else "Updated") + f" {len(updated)} item(s):")
	print("\n".join("  " + u for u in updated))
	if not_found:
		print(f"Not found ({len(not_found)}) - fix the code on the sheet or the Item:")
		print("\n".join("  " + c for c in not_found))
	if mismatched:
		print(f"Box weight does not match the Carton conversion ({len(mismatched)}):")
		print("\n".join("  " + m for m in mismatched))
	return {"updated": len(updated), "not_found": not_found, "mismatched": mismatched}
